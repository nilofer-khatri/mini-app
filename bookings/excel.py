import logging
import os
import tempfile
from pathlib import Path

from django.conf import settings
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .models import Booking

logger = logging.getLogger(__name__)

HEADERS = [
    "ID", "Booked on", "Customer", "Phone", "Date", "Time",
    "Services", "Total amount", "Payment method", "Payment status", "Status",
]
WIDTHS = [6, 22, 20, 15, 14, 12, 34, 14, 16, 16, 12]


def build_workbook(queryset=None):
    """Create an Excel workbook with one row per booking and a totals block."""
    if queryset is None:
        queryset = Booking.objects.all()
    queryset = queryset.prefetch_related("services").order_by("date", "time")

    wb = Workbook()
    ws = wb.active
    ws.title = "Bookings"
    ws.append(HEADERS)

    header_fill = PatternFill("solid", fgColor="DB2777")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"

    for i, width in enumerate(WIDTHS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = width

    for b in queryset:
        ws.append([
            b.pk,
            timezone.localtime(b.created_at).replace(tzinfo=None),
            b.customer_name,
            b.customer_phone,
            b.date,
            b.time,
            ", ".join(s.name for s in b.services.all()),
            float(b.total_amount),
            b.get_payment_method_display(),
            b.get_payment_status_display(),
            b.get_status_display(),
        ])
        row = ws.max_row
        ws.cell(row=row, column=2).number_format = "DD-MMM-YYYY hh:mm AM/PM"
        ws.cell(row=row, column=5).number_format = "DD-MMM-YYYY"
        ws.cell(row=row, column=6).number_format = "hh:mm AM/PM"
        ws.cell(row=row, column=8).number_format = "#,##0.00"

    last = ws.max_row
    if last >= 2:
        # Totals block (cancelled bookings are not counted)
        totals = [
            ("Total amount", f'=SUMIFS(H2:H{last},K2:K{last},"<>Cancelled")'),
            ("Amount received (paid)",
             f'=SUMIFS(H2:H{last},J2:J{last},"Paid",K2:K{last},"<>Cancelled")'),
            ("Amount pending (unpaid)",
             f'=SUMIFS(H2:H{last},J2:J{last},"Unpaid",K2:K{last},"<>Cancelled")'),
        ]
        row = last + 2
        for label, formula in totals:
            ws.cell(row=row, column=7, value=label).font = Font(bold=True)
            cell = ws.cell(row=row, column=8, value=formula)
            cell.font = Font(bold=True)
            cell.number_format = "#,##0.00"
            row += 1

    return wb


def export_bookings_to_excel():
    """Rebuild the Excel file from the database.

    Never raises, so a locked file (for example, open in Excel) cannot break a booking.
    """
    path = Path(settings.EXCEL_EXPORT_PATH)
    tmp = None
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        wb = build_workbook()
        fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=path.parent)
        os.close(fd)
        wb.save(tmp)
        os.replace(tmp, path)
        return True
    except Exception:
        logger.exception("Could not update the bookings Excel file")
        if tmp and os.path.exists(tmp):
            os.remove(tmp)
        return False